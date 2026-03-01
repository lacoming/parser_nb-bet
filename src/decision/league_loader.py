"""League settings loader from leagues.xlsx (openpyxl).

Expected Excel format (4 columns):
  A: Вид ставки (bet type)   — marks start of a new strategy group
  B: Лиги (league name)      — one per row, inherits group from column A
  C: условие (condition text) — on the first row of the group
  D: ROI % (e.g. 15 or 15%)  — ROI for ratio calculation, on first row of group

Condition examples:
  "Кф1 > или = 1,5, Кф1 < Кф2"
  "Кф1 > кф2, Кф2 > или = 1,4"
  "Кф1 > или = 1,5, кф1 < или равно 8"
"""
from __future__ import annotations

import logging
import re
from pathlib import Path
from typing import Optional

from src.decision.models import LeagueSetting, normalize_bet_type

log = logging.getLogger("parser_nb_bet.decision.league_loader")


def find_leagues_xlsx(search_dir: str) -> Optional[str]:
    """Scan *search_dir* for the first ``*.xlsx`` file that looks like a leagues file.

    Excludes known output/result files produced by the app itself.
    Returns the full path or ``None``.
    """
    _EXCLUDE_PREFIXES = ("result", "output", "~$")

    d = Path(search_dir)
    if not d.is_dir():
        log.debug("find_leagues_xlsx: directory does not exist: %s", search_dir)
        return None

    for p in sorted(d.glob("*.xlsx")):
        name_lower = p.name.lower()
        if any(name_lower.startswith(pref) for pref in _EXCLUDE_PREFIXES):
            continue
        # C2 fix: exclude monthly result files like 2026-03_results.xlsx
        if re.match(r"^\d{4}-\d{2}_", p.name):
            continue
        log.info("Auto-discovered leagues xlsx: %s", p)
        return str(p)

    log.debug("find_leagues_xlsx: no xlsx files found in %s", search_dir)
    return None


def load_league_settings(path: str = "leagues.xlsx") -> list[LeagueSetting]:
    """Load strategy settings from leagues.xlsx."""
    p = Path(path)
    if not p.exists():
        log.warning("leagues.xlsx not found at %s, returning empty settings", p)
        return []

    try:
        from openpyxl import load_workbook
    except ImportError:
        log.error("openpyxl not installed, cannot read leagues.xlsx")
        return []

    wb = load_workbook(filename=str(p), data_only=True)
    ws = wb.active
    if ws is None:
        log.warning("No active sheet in leagues.xlsx")
        return []

    settings: list[LeagueSetting] = []
    current: Optional[LeagueSetting] = None

    for row in range(2, (ws.max_row or 1) + 1):
        bet_type_cell = ws.cell(row=row, column=1).value
        league_cell = ws.cell(row=row, column=2).value
        condition_cell = ws.cell(row=row, column=3).value
        roi_cell = ws.cell(row=row, column=4).value

        # New group starts when column A has a value
        if bet_type_cell is not None and str(bet_type_cell).strip():
            raw_type = str(bet_type_cell).strip()
            bet_type = normalize_bet_type(raw_type)

            condition_raw = ""
            if condition_cell is not None and str(condition_cell).strip():
                condition_raw = str(condition_cell).strip()

            roi = _parse_roi(roi_cell)

            current = LeagueSetting(
                bet_type=bet_type,
                condition_raw=condition_raw,
                roi=roi,
            )
            _apply_parsed_conditions(current, condition_raw)
            settings.append(current)

        if current is None:
            continue

        # Add league name
        if league_cell is not None:
            league = str(league_cell).strip()
            if league and league != "None":
                current.leagues.append(league)

    wb.close()
    log.info("Loaded %d strategy groups from %s", len(settings), path)
    for s in settings:
        log.debug(
            "  %s: %d leagues, conditions: %s",
            s.bet_type,
            len(s.leagues),
            s.condition_raw or "(none)",
        )
    return settings


# ── Condition parser ─────────────────────────────────────────────────────


def _apply_parsed_conditions(setting: LeagueSetting, text: str) -> None:
    """Parse condition text and set bounds/relation on the setting."""
    if not text:
        return

    # Split on ", " followed by К/к (the start of next sub-condition).
    # This avoids splitting on commas inside numbers like "1,5".
    parts = re.split(r",\s*(?=[Кк])", text)

    for part in parts:
        _apply_one_condition(setting, part.strip())


def _apply_one_condition(setting: LeagueSetting, text: str) -> None:
    """Parse and apply a single sub-condition like 'Кф1 >= 1,5' or 'Кф1 < Кф2'."""
    t = text.strip().lower()
    if not t:
        return

    # Identify left variable
    left: Optional[str] = None
    if t.startswith("кф1"):
        left = "kf1"
        t = t[3:].strip()
    elif t.startswith("кф2"):
        left = "kf2"
        t = t[3:].strip()
    else:
        log.debug("Cannot parse condition (unknown left var): %s", text)
        return

    # Identify operator
    op, t = _extract_operator(t)
    if op is None:
        log.debug("Cannot parse condition (unknown operator): %s", text)
        return

    t = t.strip()

    # Right side: variable or number?
    if t.startswith("кф1") or t.startswith("кф2"):
        right_var = "kf1" if t.startswith("кф1") else "kf2"
        _apply_relation(setting, left, op, right_var)
    else:
        val = _parse_number(t)
        if val is not None:
            _apply_bound(setting, left, op, val)
        else:
            log.debug("Cannot parse condition (bad number): %s", text)


def _extract_operator(text: str) -> tuple[Optional[str], str]:
    """Extract comparison operator from start of text.

    Handles Russian forms like '> или =', '< или равно'.
    Returns (operator, remaining_text).
    """
    patterns = [
        (r"^>\s*или\s*=\s*", ">="),
        (r"^>=\s*", ">="),
        (r"^<\s*или\s*равно\s*", "<="),
        (r"^<=\s*", "<="),
        (r"^>\s*", ">"),
        (r"^<\s*", "<"),
    ]
    for pat, op in patterns:
        m = re.match(pat, text)
        if m:
            return op, text[m.end():]
    return None, text


def _apply_relation(
    setting: LeagueSetting, left: str, op: str, right_var: str,
) -> None:
    """Set kf_relation based on a variable-to-variable comparison."""
    # Normalize to canonical form: "kf1>kf2" or "kf1<kf2"
    if left == "kf1" and right_var == "kf2":
        if op in (">", ">="):
            setting.kf_relation = "kf1>kf2"
        elif op in ("<", "<="):
            setting.kf_relation = "kf1<kf2"
    elif left == "kf2" and right_var == "kf1":
        if op in (">", ">="):
            setting.kf_relation = "kf1<kf2"  # kf2 > kf1 means kf1 < kf2
        elif op in ("<", "<="):
            setting.kf_relation = "kf1>kf2"


def _apply_bound(
    setting: LeagueSetting, var: str, op: str, val: float,
) -> None:
    """Set min/max bound on the appropriate kf variable."""
    if var == "kf1":
        if op in (">=", ">"):
            setting.min_kf1 = val
        elif op in ("<=", "<"):
            setting.max_kf1 = val
    elif var == "kf2":
        if op in (">=", ">"):
            setting.min_kf2 = val
        elif op in ("<=", "<"):
            setting.max_kf2 = val


def _parse_number(text: str) -> Optional[float]:
    """Parse a number from text, handling comma as decimal separator."""
    text = text.strip()
    if not text:
        return None
    try:
        return float(text.replace(",", "."))
    except (ValueError, TypeError):
        return None


def _parse_roi(cell_value: object) -> float:
    """Parse ROI from Excel cell value.

    Accepts:
      - numeric: 15 → 0.15 (percentage), 0.15 → 0.15 (fraction)
      - string: "15%", "15", "0.15"

    Returns ROI as a fraction (e.g. 0.15 for 15%).
    """
    if cell_value is None:
        return 0.0

    if isinstance(cell_value, (int, float)):
        val = float(cell_value)
        # If > 1, treat as percentage (e.g. 15 → 0.15)
        return val / 100.0 if val > 1.0 else val

    text = str(cell_value).strip().replace(",", ".").replace("%", "")
    if not text:
        return 0.0
    try:
        val = float(text)
        return val / 100.0 if val > 1.0 else val
    except (ValueError, TypeError):
        log.debug("Cannot parse ROI value: %r", cell_value)
        return 0.0
